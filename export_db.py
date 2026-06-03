import json
import csv
import os

os.chdir('backend')

# Загружаем data.json
with open('data.json', 'r', encoding='utf-8') as f:
    data = json.load(f)

universities = data.get('universities', [])

# Экспортируем в CSV
with open('../universities.csv', 'w', newline='', encoding='utf-8-sig') as f:
    if universities:
        fieldnames = universities[0].keys()
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for uni in universities:
            # Конвертируем списки в строки для CSV
            row = uni.copy()
            for key in ['languages', 'accreditations', 'specialty_ids']:
                if key in row and isinstance(row[key], list):
                    row[key] = ';'.join(map(str, row[key]))
            writer.writerow(row)

# Сохраняем JSON в красивом формате
with open('../universities.json', 'w', encoding='utf-8') as f:
    json.dump(universities, f, ensure_ascii=False, indent=2)

# Создаем Excel
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Университеты"
    
    # Заголовки
    headers = list(universities[0].keys())
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col)
        cell.value = header
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.font = Font(bold=True, color="FFFFFF")
    
    # Данные
    for row_idx, uni in enumerate(universities, 2):
        for col, header in enumerate(headers, 1):
            val = uni.get(header)
            if isinstance(val, list):
                val = '; '.join(map(str, val))
            ws.cell(row=row_idx, column=col).value = val
    
    # Автоширина
    for col in ws.columns:
        max_length = 0
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 50)
    
    wb.save('../universities.xlsx')
    print("✓ Excel создан: universities.xlsx")
except ImportError:
    print("⚠ Excel не создан (openpyxl не установлен)")

print(f"✓ CSV создан: universities.csv")
print(f"✓ JSON создан: universities.json")
print(f"✓ Всего университетов: {len(universities)}")
